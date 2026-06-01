"""Sample tracker daily-update I/O.

Αθανάσιος (Sakis) exports completion data (with region & urbanity) from Survey
Solutions, drops the numbers into the daily template, and uploads the file here.
This module builds that template and parses an upload back into per-stratum
updates — supporting both the per-stratum template and a raw per-interview
export (one completed interview per row), which it aggregates automatically.
"""

from __future__ import annotations

import io
import unicodedata

import pandas as pd

TEMPLATE_SHEET = "Daily Update"
COLUMNS = ["ID", "Region (NUTS-2)", "Urban/Rural", "Target",
           "Completes", "Refusals", "Calls placed"]


# --------------------------------------------------------------- normalising
def _norm(s) -> str:
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.strip().lower()


def _urban_rural(value) -> str | None:
    """Map any urban/rural label to one of the two canonical Greek strings."""
    n = _norm(value)
    if not n:
        return None
    # check rural first: "ημιαστικα" also contains the urban stem "αστ".
    if any(t in n for t in ("agr", "ημι", "αγρ", "rural", "semi")) or n in ("r",):
        return "Αγροτικά/Ημιαστικά"
    if "αστ" in n or "ast" in n or "urban" in n or n in ("u",):
        return "Αστικά"
    return None


def _lookups(state: dict):
    """id->row, and normalised (region,type)->row, from current sample."""
    by_id, by_rt = {}, {}
    for r in state.get("sample", []):
        by_id[int(r["id"])] = r
        by_rt[(_norm(r["region"]), _norm(r["type"]))] = r
    return by_id, by_rt


# ------------------------------------------------------------------ template
def build_template(state: dict) -> bytes:
    """A pre-filled .xlsx (Instructions + data sheet) for Sakis to populate."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    info = wb.active
    info.title = "Instructions"
    lines = [
        ("OECD/INFE 2026 — Greece · Sample daily update", True),
        ("", False),
        ("Owner: Αθανάσιος (Sakis). Deliver every fieldwork day by 10:00.", False),
        ("", False),
        ("1. In Survey Solutions, export completed interviews with region and "
         "urban/rural.", False),
        ("2. Enter the CUMULATIVE totals to date per stratum on the "
         f"'{TEMPLATE_SHEET}' tab.", False),
        ("3. Fill 'Completes' (required). 'Refusals' and 'Calls placed' are "
         "optional.", False),
        ("4. Do NOT change the ID, Region or Urban/Rural columns — they are how "
         "the app matches rows.", False),
        ("5. Save and upload the file on the Sample page of the tracker.", False),
        ("", False),
        ("Numbers are cumulative snapshots: each upload REPLACES the previous "
         "values.", False),
    ]
    for i, (txt, bold) in enumerate(lines, 1):
        c = info.cell(row=i, column=1, value=txt)
        if bold:
            c.font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 90

    ws = wb.create_sheet(TEMPLATE_SHEET)
    head_fill = PatternFill("solid", fgColor="1A73E8")
    for j, name in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
    for i, r in enumerate(state.get("sample", []), start=2):
        ws.cell(row=i, column=1, value=int(r["id"]))
        ws.cell(row=i, column=2, value=r["region"])
        ws.cell(row=i, column=3, value=r["type"])
        ws.cell(row=i, column=4, value=int(r.get("target", 0)))
        # leave Completes / Refusals / Calls blank for Sakis
    widths = [6, 34, 22, 10, 12, 11, 13]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -------------------------------------------------------------------- parsing
def _read_any(upload) -> pd.DataFrame:
    name = getattr(upload, "name", "").lower()
    if name.endswith(".csv"):
        return pd.read_csv(upload)
    # prefer the named template sheet, else the first sheet
    xls = pd.ExcelFile(upload, engine="openpyxl")
    sheet = TEMPLATE_SHEET if TEMPLATE_SHEET in xls.sheet_names else xls.sheet_names[0]
    return xls.parse(sheet)


def _colmap(df: pd.DataFrame) -> dict:
    out = {}
    for col in df.columns:
        n = _norm(col)
        if n == "id" or n.startswith("id "):
            out["id"] = col
        elif "region" in n or "nuts" in n:
            out["region"] = col
        elif "urban" in n or "rural" in n or "αστ" in n or "αγρ" in n:
            out["type"] = col
        elif n == "target":
            out["target"] = col
        elif "complet" in n:
            out["completes"] = col
        elif "refus" in n:
            out["refusals"] = col
        elif "call" in n:
            out["calls"] = col
    return out


def parse_upload(upload, state: dict) -> dict:
    """Return {updates, mode, warnings}. 'updates' is a list of per-stratum dicts
    with old/new values for preview; nothing is written here."""
    df = _read_any(upload)
    df = df.dropna(how="all")
    cm = _colmap(df)
    by_id, by_rt = _lookups(state)
    warnings: list[str] = []

    def to_int(v):
        try:
            if pd.isna(v):
                return None
            return int(round(float(v)))
        except (TypeError, ValueError):
            return None

    # decide mode: aggregated (has a completes column) vs raw per-interview
    if "completes" in cm:
        mode = "per-stratum"
        agg: dict[int, dict] = {}
        for _, row in df.iterrows():
            target_row = None
            if "id" in cm and to_int(row[cm["id"]]) in by_id:
                target_row = by_id[to_int(row[cm["id"]])]
            elif "region" in cm and "type" in cm:
                key = (_norm(row[cm["region"]]), _norm(_urban_rural(row[cm["type"]])))
                target_row = by_rt.get(key)
            if not target_row:
                warnings.append(f"Row not matched to a stratum: "
                                f"{row.get(cm.get('region'), '?')} / "
                                f"{row.get(cm.get('type'), '?')}")
                continue
            d = agg.setdefault(int(target_row["id"]), {})
            c = to_int(row[cm["completes"]])
            if c is not None:
                d["completes"] = c
            if "refusals" in cm and to_int(row[cm["refusals"]]) is not None:
                d["refusals"] = to_int(row[cm["refusals"]])
            if "calls" in cm and to_int(row[cm["calls"]]) is not None:
                d["calls"] = to_int(row[cm["calls"]])
    else:
        # raw: one completed interview per row -> count by region × urbanity
        if "region" not in cm or "type" not in cm:
            raise ValueError("Could not find the columns to read. Expected a "
                             "'Completes' column (template), or 'Region' + "
                             "'Urban/Rural' columns (raw export).")
        mode = "raw (counted per stratum)"
        agg = {}
        for _, row in df.iterrows():
            key = (_norm(row[cm["region"]]), _norm(_urban_rural(row[cm["type"]])))
            tr = by_rt.get(key)
            if not tr:
                warnings.append(f"Interview not matched: {row[cm['region']]} / "
                                f"{row[cm['type']]}")
                continue
            d = agg.setdefault(int(tr["id"]), {})
            d["completes"] = d.get("completes", 0) + 1

    # build preview updates
    updates = []
    for sid, vals in sorted(agg.items()):
        cur = by_id[sid]
        upd = {"id": sid, "region": cur["region"], "type": cur["type"],
               "old_completes": cur.get("completes", 0),
               "new_completes": vals.get("completes", cur.get("completes", 0)),
               "new_refusals": vals.get("refusals"),
               "new_calls": vals.get("calls")}
        updates.append(upd)
    return {"updates": updates, "mode": mode, "warnings": warnings}


def apply_updates(state: dict, updates: list[dict]) -> None:
    by_id, _ = _lookups(state)
    for u in updates:
        row = by_id[u["id"]]
        row["completes"] = int(u["new_completes"])
        if u.get("new_refusals") is not None:
            row["refusals"] = int(u["new_refusals"])
        if u.get("new_calls") is not None:
            row["calls"] = int(u["new_calls"])
